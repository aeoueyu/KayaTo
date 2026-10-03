from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "KayaTo_Full-Stack_Implementation.docx"
DIAGRAM = ROOT / ".codex_tmp" / "kayato_architecture.png"

RUBY = "70000E"
RUBY_LIGHT = "F4E8EA"
INK = "2C2929"
MUTED = "666163"
CREAM = "F5F4F2"
GRAY = "E3DFDD"
LIGHT = "FAF9F8"
WHITE = "FFFFFF"
GREEN = "2D6A4F"
AMBER = "8A5A00"
RED = "9B1C1C"


def rgb(hex_string):
    return RGBColor.from_string(hex_string)


def set_run_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7D2D0", size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_bottom_border(paragraph, color=RUBY, size=16, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Title": (30, RUBY, 0, 8),
        "Subtitle": (14, MUTED, 0, 18),
        "Heading 1": (16, RUBY, 18, 10),
        "Heading 2": (13, RUBY, 14, 7),
        "Heading 3": (12, INK, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.05

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.font.color.rgb = rgb(INK)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.keep_together = True
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F3F1F0")
    code._element.get_or_add_pPr().append(shd)


def configure_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("KayaTo | Full-Stack Implementation Specification")
    set_run_font(r, size=8.5, color=MUTED, bold=True)
    add_bottom_border(p, color="D8D2D0", size=6, space=4)
    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_para(doc, text="", size=None, color=None, bold=False, italic=False,
             align=None, before=0, after=6, keep=False, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color or INK, bold=bold, italic=italic)
    return p


def add_rich_para(doc, chunks, after=6, before=0, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    for text, opts in chunks:
        r = p.add_run(text)
        set_run_font(r, size=opts.get("size"), color=opts.get("color", INK),
                     bold=opts.get("bold"), italic=opts.get("italic"))
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.keep_together = True
    return p


def numbered(doc, text):
    p = doc.add_paragraph(text, style="List Number")
    p.paragraph_format.keep_together = True
    return p


def callout(doc, label, text, tone="ruby"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent=120)
    set_table_borders(table, color=RUBY if tone == "ruby" else "CFC9C6", size=8)
    mark_header_row(table.rows[0])
    cell = table.cell(0, 0)
    set_cell_shading(cell, RUBY_LIGHT if tone == "ruby" else LIGHT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, color=RUBY if tone == "ruby" else INK, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    add_para(doc, "", after=2)


def add_table(doc, headers, rows, widths, font_size=9.4, header_fill=RUBY_LIGHT):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths, indent=120)
    set_table_borders(table)
    header = table.rows[0]
    keep_row_together(header)
    mark_header_row(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, color=RUBY, bold=True)
    for row_values in rows:
        row = table.add_row()
        keep_row_together(row)
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    add_para(doc, "", after=2)
    return table


def page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def make_architecture_diagram(path):
    w, h = 1500, 760
    img = Image.new("RGB", (w, h), "#FAF9F8")
    d = ImageDraw.Draw(img)
    try:
        font_b = ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf", 32)
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 24)
        small = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 20)
    except OSError:
        font_b = font = small = ImageFont.load_default()

    def box(x1, y1, x2, y2, title, lines, fill="#FFFFFF", outline="#70000E"):
        d.rounded_rectangle((x1, y1, x2, y2), radius=22, fill=fill, outline=outline, width=4)
        d.text((x1 + 24, y1 + 20), title, font=font_b, fill="#70000E")
        y = y1 + 70
        for line in lines:
            d.text((x1 + 24, y), line, font=small, fill="#2C2929")
            y += 31

    def arrow(x1, y1, x2, y2):
        d.line((x1, y1, x2, y2), fill="#827A7D", width=5)
        if x2 >= x1:
            pts = [(x2, y2), (x2 - 18, y2 - 10), (x2 - 18, y2 + 10)]
        else:
            pts = [(x2, y2), (x2 + 18, y2 - 10), (x2 + 18, y2 + 10)]
        d.polygon(pts, fill="#827A7D")

    box(50, 75, 380, 260, "Clients", ["React web application", "Flutter mobile application"], fill="#F4E8EA")
    box(560, 75, 930, 260, "API Gateway", ["Express REST API", "JWT / RBAC", "Socket.IO realtime"], fill="#FFFFFF")
    box(1120, 75, 1450, 260, "Application Services", ["Task and team service", "Bill and reminder service"], fill="#F4E8EA")
    box(50, 455, 380, 650, "Async Workers", ["Notifications", "File processing", "AI plan generation"], fill="#FFFFFF")
    box(560, 455, 930, 650, "Data Layer", ["MongoDB", "Redis / job queue", "Object storage"], fill="#F4E8EA")
    box(1120, 455, 1450, 650, "External Services", ["AI provider adapter", "Email / push provider", "Official biller sites"], fill="#FFFFFF")
    arrow(380, 165, 560, 165)
    arrow(930, 165, 1120, 165)
    arrow(745, 260, 745, 455)
    arrow(380, 550, 560, 550)
    arrow(930, 550, 1120, 550)
    arrow(1285, 260, 1285, 455)
    d.text((50, 700), "Recommended production topology: stateless API, durable data stores, and isolated background processing.", font=font, fill="#666163")
    img.save(path)


doc = Document()
style_document(doc)
for section in doc.sections:
    configure_header_footer(section)

# Cover
add_para(doc, "KayaTo", size=16, color=RUBY, bold=True, after=42)
add_para(doc, "FULL-STACK", size=12, color=RUBY, bold=True, after=4)
title = add_para(doc, "Implementation Specification", size=30, color=INK, bold=True, after=10)
add_bottom_border(title, color=RUBY, size=18, space=8)
add_para(doc, "Production blueprint for the web, API, realtime collaboration, AI planner, bills, notifications, and mobile clients.", size=14, color=MUTED, after=36)

meta_rows = [
    ("Document", "Technical implementation specification"),
    ("Version", "1.0"),
    ("Prepared", "27 September 2026"),
    ("Repository scope", "kayato-web and kayato-mobile"),
    ("Audience", "Product, frontend, backend, mobile, QA, and DevOps"),
    ("Status", "Implementation blueprint"),
]
add_table(doc, ["Field", "Value"], meta_rows, [2300, 7060], font_size=10.2, header_fill=CREAM)
add_para(doc, "", after=36)
callout(doc, "Purpose", "Translate the existing KayaTo prototype and API foundation into a secure, testable, deployable full-stack product. Proposed components are clearly separated from features already present in the repository.")
add_para(doc, "Prepared from the local project snapshot. No claim in this document means a proposed production feature is already implemented.", size=9.5, color=MUTED, italic=True, after=0)

page_break(doc)

# Contents and executive summary
doc.add_heading("Document map", level=1)
contents = [
    "1. Executive summary and implementation principles",
    "2. Repository baseline and gap assessment",
    "3. Target architecture and system boundaries",
    "4. Domain model and persistence design",
    "5. Backend services and API contract",
    "6. Web and mobile client implementation",
    "7. Realtime collaboration and AI planner",
    "8. Security, privacy, and resilience",
    "9. Delivery, testing, observability, and operations",
    "10. Phased roadmap and definition of done",
    "Appendices: environment, repository structure, and acceptance checklist",
]
for item in contents:
    bullet(doc, item)

doc.add_heading("1. Executive summary", level=1)
add_para(doc, "KayaTo is a productivity platform that combines personal tasks, team collaboration, AI-assisted planning, chat, and bill reminders. The repository already demonstrates the intended user experience and contains a useful MERN API foundation. The next implementation step is to replace mock state with authenticated service calls, complete the missing collaboration and AI workflows, and introduce production controls for security, reliability, and operations.")
callout(doc, "Recommended approach", "Retain the existing React, Express, MongoDB, Mongoose, JWT, Socket.IO, and Flutter choices. Add a shared API contract, service/repository separation, durable file storage, background jobs, validation, automated tests, and deployment automation before expanding feature scope.")

doc.add_heading("Implementation principles", level=2)
for text in [
    "One source of truth: server data is authoritative; client mock data is removed once an endpoint is connected.",
    "Incremental delivery: each vertical slice includes UI, API, authorization, persistence, tests, telemetry, and documentation.",
    "Secure by default: least-privilege access, validated input, short-lived sessions, safe file handling, and no payment credential storage.",
    "Human-controlled AI: Kaya proposes structured plans; users review and confirm before tasks are created or assigned.",
    "Mobile parity by API, not duplication: web and Flutter consume the same versioned backend contract.",
]:
    bullet(doc, text)

doc.add_heading("Success criteria", level=2)
add_table(doc, ["Area", "Release criterion"], [
    ("Product", "Core task, team, chat, bill, and AI-plan journeys work without mock data."),
    ("Security", "Authorization is enforced server-side for every protected resource and realtime room."),
    ("Quality", "Automated unit, integration, and end-to-end suites gate production releases."),
    ("Operations", "Health, logs, metrics, alerts, backups, and rollback procedures are documented and tested."),
    ("Experience", "Responsive web and mobile flows meet accessibility and performance budgets."),
], [1800, 7560])

# Baseline
doc.add_heading("2. Repository baseline and gap assessment", level=1)
add_para(doc, "The following baseline reflects the files inspected in the local workspace on 27 September 2026.")

doc.add_heading("Current implementation", level=2)
add_table(doc, ["Layer", "Present today", "Implementation status"], [
    ("Web client", "React 19, Vite 7, React Router, MUI, Axios, Lucide; landing, login, onboarding, dashboard, tasks, AI planner, chat, and bills screens.", "Functional prototype; most feature data and actions are local/mock."),
    ("API", "Express 5 server, CORS, JSON parsing, health endpoint, JWT middleware, MongoDB connection.", "Runnable foundation."),
    ("Authentication", "Register, login, current-user endpoint; bcrypt password hash and seven-day JWT.", "Basic implementation; production session lifecycle is incomplete."),
    ("Tasks", "Authenticated list/create/update/delete with paging and status/type filtering.", "API exists; frontend is not wired."),
    ("Bills", "Authenticated list/create/update/delete; HTTPS payment URL validation.", "API exists; frontend is not wired."),
    ("Realtime chat", "Socket.IO connection with join and send events; Message model exists.", "Transport skeleton only; persistence and room authorization are missing."),
    ("Teams", "Team and membership schemas with roles and skills.", "Data model only; no team routes or invitation lifecycle."),
    ("Mobile", "Flutter Material 3 shell and reusable toast system showcase.", "Prototype; no product navigation, API client, or auth state."),
], [1250, 4750, 3360], font_size=8.8)

doc.add_heading("Priority gaps", level=2)
for text in [
    "The React pages use mock data and local component state rather than the Express API.",
    "Authentication lacks refresh-token rotation, logout/revocation, email verification, password reset, and route guards.",
    "Team membership, invitations, conversations, message persistence, notifications, calendar, and settings APIs are absent.",
    "AI planning is a UI simulation; there is no upload, extraction, model call, schema validation, review draft, or task-creation transaction.",
    "There is no shared request/response validation, API versioning, rate limiting, audit trail, automated test suite, CI/CD pipeline, or production observability.",
    "The Flutter app is not yet a KayaTo client; it currently demonstrates notification/toast behavior only.",
]:
    bullet(doc, text)

doc.add_heading("Scope boundaries", level=2)
add_rich_para(doc, [
    ("In scope: ", {"bold": True, "color": RUBY}),
    ("task and team collaboration, authenticated chat, AI document-to-plan generation, bill reminders with safe external payment redirection, notifications, calendar aggregation, and shared web/mobile APIs.", {}),
])
add_rich_para(doc, [
    ("Out of scope for the first production release: ", {"bold": True, "color": RUBY}),
    ("processing card or bank credentials, autonomous AI task assignment without confirmation, payroll, accounting, and arbitrary third-party biller integrations.", {}),
])

# Architecture
doc.add_heading("3. Target architecture and system boundaries", level=1)
make_architecture_diagram(DIAGRAM)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
shape = p.add_run().add_picture(str(DIAGRAM), width=Inches(6.45))
doc_pr = shape._inline.docPr
doc_pr.set("descr", "KayaTo architecture diagram showing web and mobile clients, API gateway, application services, workers, data layer, and external services.")
doc_pr.set("title", "KayaTo production architecture")
add_para(doc, "Figure 1. Recommended KayaTo production architecture.", size=9, color=MUTED, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

doc.add_heading("Component responsibilities", level=2)
add_table(doc, ["Component", "Responsibility", "Key rule"], [
    ("React web", "Browser UX, routing, accessible interaction, query cache, optimistic updates.", "Never enforce ownership only in the browser."),
    ("Flutter mobile", "Native task, chat, bill, notification, and offline-aware experiences.", "Consume the same versioned API contract."),
    ("Express API", "Authentication, authorization, validation, domain orchestration, REST and Socket.IO entry points.", "Keep route handlers thin and testable."),
    ("MongoDB", "Primary durable store for users, teams, tasks, conversations, messages, bills, and audit records.", "Indexes reflect access and retention patterns."),
    ("Redis and workers", "Queues, idempotency, rate-limit counters, reminder schedules, and realtime fan-out when scaled.", "Jobs must be retry-safe."),
    ("Object storage", "Source documents, message attachments, and bill receipts.", "Private by default; access through short-lived signed URLs."),
    ("AI adapter", "Provider-neutral extraction-to-structured-plan pipeline.", "Validate model output before persistence."),
], [1480, 4920, 2960], font_size=8.7)

doc.add_heading("Backend module layout", level=2)
add_para(doc, "A modular monolith is appropriate for the first production release. It keeps transactions, deployment, and debugging simple while preserving clear service boundaries that can be extracted later if scale requires it.")
add_para(doc, "server/src/\n  config/          environment, database, logging\n  middleware/      auth, RBAC, validation, rate limits, errors\n  modules/\n    auth/ users/ teams/ tasks/ conversations/ bills/\n    ai-plans/ notifications/ files/ audit/\n  realtime/        authenticated Socket.IO handlers\n  jobs/            queue producers and worker processors\n  shared/          errors, pagination, schemas, utilities\n  app.js           Express composition\n  server.js        process startup and graceful shutdown", style="Code Block")

# Domain model
doc.add_heading("4. Domain model and persistence design", level=1)
add_para(doc, "Existing Mongoose models should be retained where practical, then tightened with indexes, status transition rules, timestamps, soft-deletion policy where needed, and service-level authorization.")

doc.add_heading("Core collections", level=2)
add_table(doc, ["Collection", "Purpose", "Important fields / indexes"], [
    ("users", "Identity and personal preferences.", "email unique; timezone; currency; onboardingCompleted; authProviders."),
    ("refresh_sessions", "Revocable, rotated login sessions.", "user, tokenHash, expiresAt TTL, device metadata, revokedAt."),
    ("teams", "Workspace ownership and role membership.", "owner; members.user; members.role; invitation references."),
    ("tasks", "Personal and team work items.", "creator, team, assignees, status, priority, dueDate; compound access/status indexes."),
    ("conversations", "Direct or channel chat membership.", "team, type, members, name, lastMessageAt."),
    ("messages", "Persistent user and Kaya messages.", "conversation + createdAt; sender; attachments; replyTo; readBy."),
    ("bills", "Reminder metadata and payment-status records.", "user + dueDate; status; cycle; officialPaymentUrl; paidAt."),
    ("file_assets", "Private uploaded-file metadata.", "owner/team, storageKey, sha256, mimeType, scanStatus, retentionUntil."),
    ("ai_plans", "AI plan drafts and provenance.", "sourceFile, requester, inputHash, model metadata, structured draft, status."),
    ("notifications", "In-app, email, and push delivery state.", "user + readAt; channel; scheduledFor; dedupeKey."),
    ("audit_events", "Security and high-value business actions.", "actor, action, resourceType/id, requestId, timestamp, metadata."),
], [1500, 3180, 4680], font_size=8.5)

doc.add_heading("Task lifecycle", level=2)
add_para(doc, "Recommended state transitions:")
add_para(doc, "todo -> requested/assigned -> in-progress -> review -> done", style="Code Block")
for text in [
    "Personal tasks may move directly from todo to in-progress or done.",
    "Team assignments require team membership; request mode requires explicit acceptance before assigned.",
    "Only the creator, authorized manager, or designated assignee may perform allowed transitions.",
    "Completing a task sets progress to 100; reopening it restores the previous non-final progress or a configured default.",
    "Every transition emits a domain event for notifications, activity history, and realtime updates.",
]:
    bullet(doc, text)

doc.add_heading("Data consistency rules", level=2)
for text in [
    "Store money as Decimal128 and serialize it through a deliberate decimal-to-string boundary; never use binary floating point for persisted amounts.",
    "Store timestamps in UTC and render using the user's IANA timezone (Asia/Manila is the current default).",
    "Use cursor pagination for high-volume messages and activity feeds; retain page pagination for moderate task lists if the UI depends on page counts.",
    "Use idempotency keys for AI confirmation, bill status updates, invitations, and retryable write operations.",
    "Do not hard-delete audit events. Define explicit retention for files, messages, notifications, and inactive accounts.",
]:
    bullet(doc, text)

# Backend/API
doc.add_heading("5. Backend services and API contract", level=1)
doc.add_heading("API conventions", level=2)
for text in [
    "Prefix production endpoints with /api/v1; keep /api/health for platform checks.",
    "Validate params, query, and body at the edge with shared schemas; reject unknown fields on sensitive writes.",
    "Return a consistent error envelope: code, message, details, and requestId.",
    "Use ISO 8601 timestamps, stable enum values, and explicit decimal serialization.",
    "Publish an OpenAPI document and generate typed web/mobile clients from the same contract where practical.",
]:
    bullet(doc, text)

doc.add_heading("Endpoint inventory", level=2)
add_table(doc, ["Module", "Endpoints", "Status / notes"], [
    ("Health", "GET /api/health", "Existing; add readiness checks separately from liveness."),
    ("Auth", "POST /auth/register, /login; GET /auth/me", "Existing foundation. Add refresh, logout, verify-email, forgot/reset password, and session list/revoke."),
    ("Tasks", "GET/POST /tasks; PATCH/DELETE /tasks/:id", "Existing foundation. Add single-item GET, transition endpoint, comments, attachments, and team authorization."),
    ("Bills", "GET/POST /bills; PATCH/DELETE /bills/:id", "Existing foundation. Add mark-paid action, recurrence materialization, reminders, and safe URL registry."),
    ("Teams", "GET/POST /teams; GET/PATCH /teams/:id; POST /teams/:id/invitations; PATCH /members/:userId", "Required."),
    ("Chat", "GET/POST /conversations; GET/POST /conversations/:id/messages; PATCH read state", "Required; pair with authenticated Socket.IO events."),
    ("AI plans", "POST /files; POST /ai-plans; GET/PATCH /ai-plans/:id; POST /ai-plans/:id/confirm", "Required; asynchronous generation and idempotent confirmation."),
    ("Notifications", "GET /notifications; PATCH /notifications/:id/read; GET/PATCH /preferences", "Required."),
    ("Calendar", "GET /calendar?from=&to=", "Aggregate tasks, bill due dates, and reminders without duplicating source records."),
], [1160, 4250, 3950], font_size=8.2)

doc.add_heading("Service-layer authorization", level=2)
add_para(doc, "Authorization must use resource-aware policies rather than checking only for a valid token. Route middleware authenticates the actor; service policies decide whether the actor can read or mutate the resource.")
add_table(doc, ["Resource", "Read", "Write"], [
    ("Personal task", "Creator or assignee as configured", "Creator; assignee only for permitted progress/status fields"),
    ("Team task", "Active team member with sufficient role", "Owner/manager; assignee for permitted transitions"),
    ("Conversation", "Active conversation member", "Active member; channel administration by team role"),
    ("Bill", "Owning user only", "Owning user only"),
    ("AI plan", "Requester and explicitly authorized team members", "Requester until confirmed; confirmation requires task-create permission"),
], [1800, 3500, 4060], font_size=8.8)

doc.add_heading("Representative response envelope", level=2)
add_para(doc, '{\n  "data": { "id": "...", "title": "Finalize mobile onboarding" },\n  "meta": { "requestId": "req_..." }\n}\n\n{\n  "error": {\n    "code": "TASK_TRANSITION_INVALID",\n    "message": "Task cannot move from done to review.",\n    "details": {},\n    "requestId": "req_..."\n  }\n}', style="Code Block")

# Clients
doc.add_heading("6. Web and mobile client implementation", level=1)
doc.add_heading("React web application", level=2)
add_para(doc, "Keep the existing route and visual structure, but introduce explicit application layers so pages do not own transport or server-cache logic.")
add_para(doc, "src/\n  app/             router, providers, auth bootstrap\n  api/             generated/manual API client and interceptors\n  features/        auth, tasks, teams, chat, bills, ai-planner\n  components/      shared accessible UI\n  hooks/            reusable application hooks\n  state/            small client-only state; no duplicate server cache\n  theme/            tokens and light/dark modes", style="Code Block")

doc.add_heading("Web integration sequence", level=3)
for text in [
    "Add an AuthProvider that restores the current session, protects /app routes, and handles refresh/logout.",
    "Create a single Axios instance with base URL, credentials, timeout, request ID, cancellation, and normalized errors.",
    "Adopt a server-state library or equivalent cache layer for queries, mutations, invalidation, optimistic updates, and retry policy.",
    "Replace mockData usage one feature at a time; retain fixture data only in Storybook/tests.",
    "Model loading, empty, offline, error, permission-denied, and stale-data states for every page.",
    "Use semantic HTML, keyboard-accessible dialogs/boards, visible focus, reduced-motion support, and automated accessibility checks.",
]:
    numbered(doc, text)

doc.add_heading("Flutter mobile application", level=2)
add_para(doc, "The current Flutter code should be treated as a reusable theme and toast foundation. Build the application shell around feature modules and a shared API contract.")
add_table(doc, ["Area", "Recommended implementation"], [
    ("Navigation", "Authenticated shell with Dashboard, Tasks, Chat, Bills, and Profile tabs; guarded onboarding/auth routes."),
    ("State", "Predictable feature state with immutable models; separate transient UI state from remote/cache state."),
    ("Networking", "Typed client, secure token storage, refresh coordination, request cancellation, retries only for safe/idempotent operations."),
    ("Realtime", "Socket lifecycle bound to authenticated session; reconnect with backoff; reconcile with paged message history."),
    ("Offline", "Read-through cache for tasks/conversations; queued mutations only where conflicts are understood and surfaced."),
    ("Notifications", "Push token registration, deep links, and in-app toast behavior using the existing notification component."),
], [1760, 7600], font_size=9.0)

doc.add_heading("Shared UX contract", level=2)
for text in [
    "Use identical status, priority, role, and error-code enums across clients.",
    "Display all dates in the user's timezone and all amounts using the user's currency/locale.",
    "Require confirmation for destructive actions, AI plan creation, and outbound bill-payment navigation.",
    "Show optimistic state only when the operation can be safely rolled back; otherwise show a pending state.",
    "Deep links identify a resource, never trust client-provided permissions, and resolve through the API.",
]:
    bullet(doc, text)

# Realtime and AI
doc.add_heading("7. Realtime collaboration and AI planner", level=1)
doc.add_heading("Authenticated chat flow", level=2)
for text in [
    "The client connects with a valid access token; the server resolves the user and rejects unauthenticated sockets.",
    "conversation:join verifies membership before adding the socket to a room.",
    "message:send validates content and attachment ownership, persists the message, then emits the canonical stored message.",
    "message:new includes a server-generated ID, timestamps, sender summary, and correlation/request ID.",
    "Read receipts and typing indicators are rate-limited; durable read state is persisted separately from ephemeral typing state.",
    "On reconnect, the client fetches messages after its latest known cursor to close event gaps.",
]:
    numbered(doc, text)

add_table(doc, ["Event", "Direction", "Purpose"], [
    ("conversation:join", "Client -> server", "Request authorized room membership."),
    ("message:send", "Client -> server", "Validate and persist a message; acknowledge with canonical record."),
    ("message:new", "Server -> room", "Deliver persisted message to authorized members."),
    ("message:read", "Both", "Persist and fan out read position."),
    ("task:updated", "Server -> team/user", "Synchronize task changes across clients."),
    ("notification:new", "Server -> user", "Deliver in-app notifications."),
], [1800, 1900, 5660], font_size=8.8)

doc.add_heading("AI planner pipeline", level=2)
for text in [
    "Upload: validate extension, MIME signature, size, ownership, and checksum; store privately.",
    "Scan and extract: virus-scan, extract text in an isolated worker, and reject encrypted or unsupported documents with a clear error.",
    "Prepare: normalize text, enforce limits, redact configured sensitive patterns, and record a content hash.",
    "Generate: call an AI provider through an adapter that requests a strict structured schema.",
    "Validate: schema-check every proposed task, dates, estimates, priorities, and assignee references; discard unknown members.",
    "Review: present an editable draft with provenance, warnings, confidence notes, and no assignments applied yet.",
    "Confirm: create tasks idempotently, log the action, emit events, and preserve the approved draft for auditability.",
]:
    numbered(doc, text)

callout(doc, "AI control", "Kaya may recommend task structure and assignees, but the system must not create, assign, delete, or mark work complete until an authorized user explicitly confirms the operation.")

doc.add_heading("Failure behavior", level=2)
for text in [
    "Queue jobs with bounded exponential backoff and a dead-letter state; never retry permanent validation failures.",
    "Expose plan status as uploaded, processing, ready, failed, confirmed, or expired.",
    "Make confirm idempotent so a client retry cannot create duplicate tasks.",
    "Keep the original document and extracted text under an explicit retention policy; allow authorized deletion.",
    "Record provider, model identifier, prompt/template version, latency, token usage, and validation result without logging source content by default.",
]:
    bullet(doc, text)

# Security
doc.add_heading("8. Security, privacy, and resilience", level=1)
doc.add_heading("Authentication and session design", level=2)
for text in [
    "Use short-lived access tokens and rotated refresh tokens stored as hashes server-side.",
    "For web, prefer a Secure, HttpOnly, SameSite refresh cookie and keep access tokens in memory; implement CSRF protection where cookies authorize writes.",
    "For mobile, store refresh credentials only in platform secure storage and bind sessions to device metadata where appropriate.",
    "Provide logout, logout-all, active-session review, password reset, email verification, and token revocation after credential changes.",
    "Apply login throttling, generic credential errors, password-strength checks, and breached-password screening if available.",
]:
    bullet(doc, text)

doc.add_heading("Threat controls", level=2)
add_table(doc, ["Risk", "Required control"], [
    ("Broken object authorization", "Policy checks on every resource and socket room; never trust IDs or roles supplied by the client."),
    ("Injection / malformed input", "Schema validation, Mongoose strict mode, safe query construction, output encoding, and controlled filters/sorts."),
    ("File attacks", "MIME sniffing, allowlists, size limits, malware scan, isolated extraction, private storage, and signed downloads."),
    ("Account/session theft", "TLS, secure cookie/storage, refresh rotation, revocation, device/session visibility, and anomaly alerts."),
    ("Abuse and cost spikes", "Per-IP/user limits, AI quotas, upload quotas, queue concurrency limits, and budget alarms."),
    ("Sensitive logging", "Structured redaction; exclude passwords, tokens, document contents, and payment-related personal data."),
    ("Unsafe bill redirect", "Curated/verified HTTPS domains, clear outbound warning, hostname display, and no credential collection by KayaTo."),
    ("Supply-chain risk", "Lockfiles, dependency scanning, minimal runtime images, provenance, and prompt patching with tested upgrades."),
], [2350, 7010], font_size=8.8)

doc.add_heading("Privacy requirements", level=2)
for text in [
    "Publish a data inventory and purpose for each field; collect only what a feature requires.",
    "Define retention and deletion for accounts, uploaded files, extracted text, AI plans, chat messages, receipts, and audit logs.",
    "Provide export and deletion workflows with team-owned data rules clearly separated from personal data.",
    "Encrypt data in transit and at rest; restrict production access with least privilege and audited administrative actions.",
    "Do not use customer content for model training unless the user has explicitly opted in under a clear policy.",
]:
    bullet(doc, text)

doc.add_heading("Reliability patterns", level=2)
for text in [
    "Graceful shutdown stops new traffic, drains requests, closes Socket.IO, and then closes database/queue connections.",
    "Readiness fails when mandatory dependencies are unavailable; liveness verifies only that the process can respond.",
    "Use retries only for transient failures, with jitter and idempotency; add circuit breakers around unstable external services.",
    "Back up MongoDB and object storage, test restoration, and document recovery time and recovery point objectives.",
]:
    bullet(doc, text)

# Delivery & testing
doc.add_heading("9. Delivery, testing, observability, and operations", level=1)
doc.add_heading("Environment model", level=2)
add_table(doc, ["Environment", "Purpose", "Data policy"], [
    ("Local", "Developer workflow with seeded fixtures and local/ephemeral dependencies.", "Synthetic data only."),
    ("CI", "Automated lint, tests, build, migrations/index checks, and security scanning.", "Ephemeral test data."),
    ("Staging", "Production-like integration, UAT, load tests, and release rehearsal.", "Synthetic or explicitly sanitized data."),
    ("Production", "Customer traffic with monitored, least-privilege infrastructure.", "Real data under documented retention and access policies."),
], [1500, 4440, 3420], font_size=9.0)

doc.add_heading("CI/CD pipeline", level=2)
for text in [
    "Install from lockfiles and verify formatting, linting, type/static checks, and unit tests.",
    "Run API integration tests against an isolated MongoDB instance and contract tests against the OpenAPI schema.",
    "Build web, API container, Android, and iOS artifacts; generate a software bill of materials and scan dependencies/images.",
    "Deploy automatically to staging, run smoke and end-to-end tests, then promote the immutable artifact with an approval gate.",
    "Run post-deploy health checks and support immediate rollback to the previous known-good version.",
]:
    numbered(doc, text)

doc.add_heading("Test strategy", level=2)
add_table(doc, ["Level", "Coverage"], [
    ("Unit", "Domain transitions, authorization policies, validators, recurrence calculations, serializers, and AI-output parsing."),
    ("Integration", "Routes + MongoDB, session rotation, team membership, Socket.IO authorization, queues, file metadata, and idempotency."),
    ("Contract", "OpenAPI request/response compatibility for React and Flutter clients."),
    ("End to end", "Register/login, create/assign/complete task, team invitation, chat, AI review/confirm, bill reminder/mark paid."),
    ("Accessibility", "Keyboard flows, dialogs, focus order, labels, color contrast, reduced motion, and screen-reader smoke tests."),
    ("Performance", "Task list paging/filtering, message history, socket fan-out, upload/extraction queue, and dashboard aggregation."),
    ("Security", "Authorization matrix, rate limits, dependency scan, secret scan, file-abuse tests, and session-revocation tests."),
], [1600, 7760], font_size=8.8)

doc.add_heading("Observability", level=2)
for text in [
    "Structured JSON logs with requestId, actor ID where appropriate, route, status, latency, and error code.",
    "Metrics for request rate/errors/latency, database latency, socket connections, queue depth/age, AI latency/cost, notification delivery, and job failures.",
    "Distributed tracing or correlation across API, worker, database, object storage, and AI provider calls.",
    "Alerts tied to user impact: elevated error rate, authentication failures, queue backlog, reminder lateness, and AI failure/cost thresholds.",
    "Dashboards and runbooks for login failure, database degradation, realtime disconnect spikes, stalled AI jobs, and missed reminders.",
]:
    bullet(doc, text)

doc.add_heading("Suggested service objectives", level=2)
add_table(doc, ["Capability", "Initial target"], [
    ("Authenticated API", "99.9% monthly availability; p95 read latency below 400 ms excluding external calls."),
    ("Realtime message delivery", "p95 server-to-room emit below 1 second under expected load."),
    ("AI plan generation", "95% of supported documents reach ready or an actionable failed state within 2 minutes."),
    ("Bill reminders", "99% of scheduled reminders dispatched within 5 minutes of target time."),
], [2500, 6860], font_size=9.0)

# Roadmap
doc.add_heading("10. Phased roadmap and definition of done", level=1)
add_para(doc, "The sequence below prioritizes a working vertical slice before adding high-cost integrations. Durations are planning estimates for a small cross-functional team and should be recalibrated after backlog sizing.")
add_table(doc, ["Phase", "Outcome", "Key deliverables", "Indicative duration"], [
    ("0. Foundations", "Repeatable engineering baseline", "Environment validation, app composition, shared errors/validation, logging, OpenAPI, test harness, CI.", "1 week"),
    ("1. Auth + API wiring", "Real authenticated web session", "Refresh/logout, route guards, Axios client, current-user bootstrap, task and bill screens connected.", "2 weeks"),
    ("2. Teams + tasks", "Complete collaborative task flow", "Teams, invites, RBAC, assignments, transitions, comments, activity, notifications.", "2-3 weeks"),
    ("3. Chat + realtime", "Persistent authorized collaboration", "Conversations, messages, Socket.IO auth, reconnect sync, read state, attachments.", "2 weeks"),
    ("4. AI planner", "Document-to-task workflow", "Private uploads, scan/extract workers, provider adapter, validated drafts, review, idempotent confirmation.", "2-3 weeks"),
    ("5. Bills + calendar", "Reliable personal reminders", "Recurrence, scheduler, preferences, calendar aggregation, safe biller redirects, mark-paid flow.", "2 weeks"),
    ("6. Mobile + hardening", "Production release candidate", "Flutter core journeys, push/deep links, accessibility, load/security testing, observability, backup/restore, release rehearsal.", "3-4 weeks"),
], [1200, 2200, 4520, 1440], font_size=8.0)

doc.add_heading("Definition of done for each feature", level=2)
for text in [
    "User-visible acceptance criteria are met on supported web breakpoints and required mobile platforms.",
    "API input/output schema, authorization policy, indexes, audit behavior, and error cases are implemented.",
    "Loading, empty, error, offline/reconnect, permission-denied, and destructive-confirmation states are designed and tested.",
    "Unit/integration/contract/end-to-end coverage appropriate to risk is passing in CI.",
    "Logs, metrics, alerts, dashboards, and an operational runbook exist for the feature's critical path.",
    "Accessibility review, security review, privacy/retention review, and product sign-off are complete.",
    "Documentation and API contract are updated; rollback behavior is known and tested.",
]:
    bullet(doc, text)

doc.add_heading("Release gates", level=2)
add_table(doc, ["Gate", "Evidence required"], [
    ("Functional", "All critical journeys pass in staging with seeded and failure-case data."),
    ("Security", "No open critical/high findings; authorization suite and session/file abuse tests pass."),
    ("Performance", "Expected peak load passes with agreed latency/error budgets and queue headroom."),
    ("Recovery", "Backup restore and application rollback are successfully rehearsed."),
    ("Operations", "On-call owner, dashboards, alerts, escalation path, and runbooks are active."),
    ("Product", "Scope, support messaging, privacy notice, and user-facing limitations are approved."),
], [1700, 7660], font_size=9.0)

# Appendices
doc.add_heading("Appendix A. Environment configuration", level=1)
add_table(doc, ["Variable / secret", "Purpose"], [
    ("NODE_ENV", "Runtime mode and safe defaults."),
    ("PORT", "API listener port."),
    ("CLIENT_URL", "Allowed web origin; use an explicit production allowlist."),
    ("MONGODB_URI", "Primary database connection."),
    ("JWT_ACCESS_SECRET / keys", "Access-token signing; manage and rotate through a secret manager."),
    ("REFRESH_TOKEN_PEPPER", "Additional protection for stored refresh-token hashes."),
    ("REDIS_URL", "Queues, rate limits, caching, and Socket.IO scaling."),
    ("OBJECT_STORAGE_*", "Private bucket, region/endpoint, credentials, and signing settings."),
    ("AI_PROVIDER / AI_API_KEY", "Provider selection and credential; never expose to clients."),
    ("EMAIL_* / PUSH_*", "Notification provider settings and credentials."),
    ("LOG_LEVEL / ERROR_TRACKING_DSN", "Operational telemetry configuration."),
], [3100, 6260], font_size=8.9)

doc.add_heading("Appendix B. Repository implementation map", level=1)
add_table(doc, ["Existing path", "Role in the implementation"], [
    ("kayato-web/src/pages", "Current product screens; refactor into feature modules while preserving approved UI."),
    ("kayato-web/src/data/mockData.js", "Prototype fixtures; move to test/story fixtures after API wiring."),
    ("kayato-web/server/src/routes", "Existing auth/task/bill routes; migrate into versioned modules with schemas and services."),
    ("kayato-web/server/src/models", "Existing User, Team, Task, Message, and Bill schemas; evolve with indexes and policies."),
    ("kayato-web/server/src/server.js", "Current process composition; split app construction from process startup for tests."),
    ("kayato-mobile/lib/main.dart", "Current Flutter theme and toast showcase; evolve into an authenticated app shell."),
    ("kayato-mobile/lib/kayato_toast.dart", "Reusable mobile notification/toast component."),
], [3600, 5760], font_size=8.8)

doc.add_heading("Appendix C. Acceptance checklist", level=1)
check_items = [
    "[ ] Web pages no longer depend on production mock data.",
    "[ ] All protected REST and Socket.IO actions enforce resource-level authorization.",
    "[ ] Refresh, logout, revocation, reset, and verification flows are complete.",
    "[ ] Team invites, roles, task assignments, transitions, and notifications work end to end.",
    "[ ] Messages persist before broadcast and recover correctly after reconnect.",
    "[ ] AI output is schema-validated and requires explicit confirmation before creating tasks.",
    "[ ] Uploaded files are private, scanned, access-controlled, and governed by retention.",
    "[ ] Bills never collect payment credentials and redirect only with a clear safety interstitial.",
    "[ ] CI gates tests, contract compatibility, security scans, and production builds.",
    "[ ] Dashboards, alerts, backups, restore tests, rollback, and runbooks are ready.",
    "[ ] Accessibility and privacy reviews are complete for release scope.",
]
for item in check_items:
    bullet(doc, item)

doc.add_heading("Implementation decision record", level=2)
callout(doc, "Decision", "Begin as a modular monolith with one versioned API and shared domain services. Split services only when measured scaling, ownership, or deployment constraints justify the added operational complexity.", tone="neutral")

# Core properties
doc.core_properties.title = "KayaTo Full-Stack Implementation Specification"
doc.core_properties.subject = "Production implementation blueprint for KayaTo"
doc.core_properties.author = "KayaTo Project Team"
doc.core_properties.keywords = "KayaTo, full-stack, React, Express, MongoDB, Flutter, Socket.IO, AI planner"
doc.core_properties.comments = "Prepared from the local KayaTo repository snapshot."

doc.save(OUT)
print(OUT)
